from datetime import date

import pytest
from openpyxl import load_workbook

from conftest import DATA
from legal_workflows.cli import run_audit
from legal_workflows.cms import AmbiguousClaimError, ClaimNotFoundError, resolve_claim_id
from legal_workflows.medical_audit.extract import AUDIT_COLUMNS, SOURCE_COLUMNS, build_rows, field_value
from legal_workflows.medical_audit.workbook import (
    FIRST_DATA_ROW, HEADER_ROW, NO_PROVIDERS_NOTICE, SHEET_NAME, safe_filename,
)


# --- resolving the claim -------------------------------------------------------

def test_resolve_from_url(client):
    assert resolve_claim_id(client, "https://cms.example.com/claims?id=10001") == 10001


def test_resolve_from_id_and_name(client):
    assert resolve_claim_id(client, "10002") == 10002
    assert resolve_claim_id(client, "rivera") == 10001


def test_ambiguous_name_is_never_guessed(client):
    with pytest.raises(AmbiguousClaimError):
        resolve_claim_id(client, "e")  # appears in every case name


def test_unknown_claim(client):
    with pytest.raises(ClaimNotFoundError):
        resolve_claim_id(client, "Nobody v. Nothing")


# --- extracting provider rows --------------------------------------------------

def test_only_medical_providers_case_insensitive(client):
    _, rows = build_rows(client.get_claim_contacts(10001))
    assert [r["Provider"] for r in rows] == [
        "Northside Chiropractic", "Lakeview Imaging Center", "Metro Pain & Spine",
        "Riverbend Physical Therapy",
    ]


def test_column_order_and_extra_fields(client):
    columns, _ = build_rows(client.get_claim_contacts(10001))
    assert columns[: len(SOURCE_COLUMNS)] == SOURCE_COLUMNS
    assert columns[-len(AUDIT_COLUMNS):] == AUDIT_COLUMNS
    extras = columns[len(SOURCE_COLUMNS): -len(AUDIT_COLUMNS)]
    assert extras == ["Total Billed", "First Visit", "Referred By", "Visit Count"]
    assert "Empty Field Example" not in columns  # empty values are not invented


def test_field_types():
    assert field_value({"type": "date", "value": "2026-03-04"}) == date(2026, 3, 4)
    assert field_value({"type": "number", "value": 6}) == 6
    assert field_value({"type": "decimal", "value": "10.5"}) == 10.5
    assert field_value({"type": "contact_reference", "value": {"name": "X"}}) == "X"
    assert field_value({"type": "short_text", "value": ""}) is None


def test_missing_collections_are_guarded():
    assert build_rows(None)[1] == []
    assert build_rows([{"role": "Medical Provider", "ad_hoc_fields": None}])[1][0]["Provider"] is None


# --- workbook ------------------------------------------------------------------

def test_workbook_end_to_end(tmp_path):
    path = run_audit("10001", DATA, tmp_path)
    assert path.name == "Rivera v. Example Transport LLC_Medical_Provider_Audit.xlsx"

    ws = load_workbook(path)[SHEET_NAME]
    assert "Providers: 4" in ws["A3"].value
    assert ws["A4"].value.endswith("https://cms.example.com/claims?id=10001")
    assert ws.freeze_panes == f"B{FIRST_DATA_ROW}"
    assert "MedicalProviderAudit" in ws.tables

    headers = [c.value for c in ws[HEADER_ROW]]
    status_col = headers.index("Audit Status") + 1
    id_col = headers.index("Contact ID") + 1
    for r in range(FIRST_DATA_ROW, FIRST_DATA_ROW + 4):
        assert ws.cell(row=r, column=status_col).value == "Not Reviewed"
        assert isinstance(ws.cell(row=r, column=id_col).value, int)

    formulas = {dv.formula1 for dv in ws.data_validations.dataValidation}
    assert '"Not Reviewed,Reviewed - No Issue,Reviewed - Issue"' in formulas
    assert '"Yes,No"' in formulas


def test_no_providers_shows_notice(tmp_path):
    ws = load_workbook(run_audit("10002", DATA, tmp_path))[SHEET_NAME]
    assert ws.cell(row=FIRST_DATA_ROW, column=1).value == NO_PROVIDERS_NOTICE
    assert ws.cell(row=HEADER_ROW, column=1).value == "Provider"


def test_safe_filename_only_replaces_invalid_characters():
    assert safe_filename("Nguyen / Sample: Claim") == "Nguyen _ Sample_ Claim_Medical_Provider_Audit.xlsx"
