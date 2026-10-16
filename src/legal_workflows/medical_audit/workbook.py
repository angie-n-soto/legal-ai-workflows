"""Build, verify and save the Medical Provider Audit workbook."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .extract import AUDIT_COLUMNS, AUDIT_STATUS_OPTIONS

SHEET_NAME = "Medical Provider Audit"
NO_PROVIDERS_NOTICE = "No medical providers found"
HEADER_ROW = 6
FIRST_DATA_ROW = HEADER_ROW + 1
WRAP_COLUMNS = {"Notes", "Role to Case", "Issue / Discrepancy", "Audit Notes"}
MIN_WIDTH, MAX_WIDTH = 10, 45
FORMULA_ERRORS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")

NAVY = "1F3A5F"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
AUDIT_HEADER_FILL = PatternFill("solid", fgColor="7A5C1E")
THIN = Side(style="thin", color="BFC7D5")


def safe_filename(case_name: str) -> str:
    """<Case_Name>_Medical_Provider_Audit.xlsx, replacing only characters invalid in filenames."""
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", case_name).strip().rstrip(".")
    return f"{cleaned or 'Case'}_Medical_Provider_Audit.xlsx"


def build_workbook(claim: dict, columns: list[str], rows: list[dict[str, Any]]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    _write_title_band(ws, claim, len(rows), len(columns))
    _write_header(ws, columns)

    if rows:
        for r, row in enumerate(rows, start=FIRST_DATA_ROW):
            for c, column in enumerate(columns, start=1):
                cell = ws.cell(row=r, column=c, value=row.get(column))
                cell.alignment = Alignment(vertical="top", wrap_text=column in WRAP_COLUMNS)
                cell.border = Border(bottom=THIN)
                if isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
        last_row = FIRST_DATA_ROW + len(rows) - 1
        ref = f"A{HEADER_ROW}:{get_column_letter(len(columns))}{last_row}"
        table = Table(displayName="MedicalProviderAudit", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        _add_dropdowns(ws, columns, FIRST_DATA_ROW, last_row)
    else:
        ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(columns))}{HEADER_ROW}"
        ws.merge_cells(
            start_row=FIRST_DATA_ROW, start_column=1,
            end_row=FIRST_DATA_ROW, end_column=len(columns),
        )
        notice = ws.cell(row=FIRST_DATA_ROW, column=1, value=NO_PROVIDERS_NOTICE)
        notice.font = Font(bold=True, italic=True, color="9C2B2B")
        notice.alignment = Alignment(horizontal="center")

    # Header colors come after the table so they win over the table style.
    for c, column in enumerate(columns, start=1):
        cell = ws.cell(row=HEADER_ROW, column=c)
        cell.fill = AUDIT_HEADER_FILL if column in AUDIT_COLUMNS else HEADER_FILL

    _set_widths(ws, columns, rows)
    ws.freeze_panes = f"B{FIRST_DATA_ROW}"
    ws.sheet_view.zoomScale = 100
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    return wb


def _write_title_band(ws, claim: dict, provider_count: int, width: int) -> None:
    lines = [
        ("Medical Provider Audit", Font(bold=True, size=16, color=NAVY)),
        (f"Case: {claim.get('case_name', '')}", Font(bold=True, size=12)),
        (f"Claim ID: {claim.get('id', '')}    |    Providers: {provider_count}", Font(size=11)),
        (f"Source: {claim.get('url', '')}", Font(size=10, color="555555")),
    ]
    for r, (text, font) in enumerate(lines, start=1):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(width, 1))
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = font
        # Keep the URL as plain text, not a hyperlink.
        cell.hyperlink = None


def _write_header(ws, columns: list[str]) -> None:
    for c, column in enumerate(columns, start=1):
        cell = ws.cell(row=HEADER_ROW, column=c, value=column)
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[HEADER_ROW].height = 30


def _add_dropdowns(ws, columns: list[str], first: int, last: int) -> None:
    status_col = get_column_letter(columns.index("Audit Status") + 1)
    status = DataValidation(
        type="list", formula1='"' + ",".join(AUDIT_STATUS_OPTIONS) + '"', allow_blank=False
    )
    status.add(f"{status_col}{first}:{status_col}{last}")
    ws.add_data_validation(status)

    follow_col = get_column_letter(columns.index("Follow-up Needed") + 1)
    follow = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    follow.add(f"{follow_col}{first}:{follow_col}{last}")
    ws.add_data_validation(follow)


def _set_widths(ws, columns: list[str], rows: list[dict[str, Any]]) -> None:
    for c, column in enumerate(columns, start=1):
        longest = max([len(column)] + [len(str(r.get(column) or "")) for r in rows])
        ws.column_dimensions[get_column_letter(c)].width = max(MIN_WIDTH, min(MAX_WIDTH, longest + 2))


def verify_workbook(wb: Workbook, expected_providers: int) -> list[str]:
    """Pre-export checks. Returns a list of problems; an empty list means it passed."""
    problems: list[str] = []
    ws = wb[SHEET_NAME]

    if expected_providers == 0:
        if ws.cell(row=FIRST_DATA_ROW, column=1).value != NO_PROVIDERS_NOTICE:
            problems.append("Missing 'No medical providers found' notice")
    else:
        data_rows = sum(
            1 for r in range(FIRST_DATA_ROW, ws.max_row + 1) if ws.cell(row=r, column=1).value
        )
        if data_rows != expected_providers:
            problems.append(f"Provider count mismatch: workbook {data_rows}, source {expected_providers}")

    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() in FORMULA_ERRORS:
                problems.append(f"Formula error {cell.value} at {cell.coordinate}")

    for c in range(1, ws.max_column + 1):
        header = ws.cell(row=HEADER_ROW, column=c).value
        width = ws.column_dimensions[get_column_letter(c)].width
        if header and width and width > MAX_WIDTH:
            problems.append(f"Column '{header}' is too wide ({width})")

    return problems


def save_workbook(wb: Workbook, claim: dict, out_dir: str | Path) -> Path:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    path = out / safe_filename(str(claim.get("case_name", "Case")))
    wb.save(path)
    return path
