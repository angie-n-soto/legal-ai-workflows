"""Draft reduction request letters from ledger data.

Medical providers get a percentage reduction of the bill. Litigation-funding companies get a
percentage reduction of interest only, plus a request for a principal/interest breakdown.
Amounts come from the CMS ledger; anything missing or suspicious is flagged, not guessed.
"""

from __future__ import annotations

import re
import shutil
import subprocess
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path

from docx import Document

from ..common import fmt_money, money, to_date
from .templates import build_template

PROVIDER, LOAN = "provider", "loan"
LOAN_ROLES = {"litigation funding", "loan company", "funding company"}
_PLACEHOLDER = re.compile(r"\{\{[A-Z_]+\}\}")


class NeedsInputError(ValueError):
    """The request can't be drafted without the user's answer (for example, a target reduction)."""


@dataclass
class Letter:
    entity: str
    kind: str
    billed: Decimal
    payoff: Decimal | None
    reduction_percent: Decimal
    filename: str
    values: dict[str, str]
    flags: list[str] = field(default_factory=list)
    docx_path: Path | None = None
    pdf_path: Path | None = None


def letter_kind(contact: dict) -> str:
    return LOAN if str(contact.get("role", "")).casefold() in LOAN_ROLES else PROVIDER


def output_filename(entity: str, client_name: str) -> str:
    """Reduction Request - Provider Name - Client Name.pdf, nothing appended."""
    name = f"Reduction Request - {entity} - {client_name}"
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", name).strip() + ".pdf"


def _pct(value) -> str:
    text = f"{Decimal(str(value)).normalize():f}"
    return f"{text}%"


def _ledger_entry(ledger: list[dict], contact_id: int, ledger_id: int | None = None) -> dict:
    entries = [e for e in ledger if e.get("contact_id") == contact_id
               and (ledger_id is None or e.get("ledger_id") == ledger_id)]
    if not entries:
        raise NeedsInputError(f"No ledger entry for contact {contact_id}")
    if len(entries) > 1:
        raise NeedsInputError(f"Contact {contact_id} has {len(entries)} ledger entries; add 'ledger_id' to the request")
    return entries[0]


def plan_letter(claim: dict, contact: dict, ledger: list[dict], request: dict,
                as_of: date, attorney_name: str) -> Letter:
    client = claim["client"]
    client_name = f"{client['first_name']} {client['last_name']}"
    entry = _ledger_entry(ledger, contact["contact_id"], request.get("ledger_id"))
    kind = letter_kind(contact)
    flags: list[str] = []
    billed = money(entry["charge"])

    address = contact.get("address") or ""
    other = contact.get("billing_address")
    if other and address and other.strip() != address.strip():
        flags.append(f"Address conflict: CMS has '{address}' and billing address '{other}'. "
                     "Used the main address; confirm before sending.")
    if not address:
        raise NeedsInputError(f"No mailing address for {contact['name']}")

    values = {
        "LETTER_DATE": f"{as_of:%B} {as_of.day}, {as_of.year}",
        "ENTITY_NAME": contact["name"],
        "ENTITY_ADDRESS": address,
        "CLIENT_NAME": client_name,
        "DOB": f"{to_date(client['dob']):%m/%d/%Y}",
        "DOI": f"{to_date(claim['date_of_incident']):%m/%d/%Y}",
        "INCIDENT_TYPE": claim.get("incident_type", "incident"),
        "ATTORNEY_NAME": attorney_name,
        "BILLED": fmt_money(billed),
    }

    if kind == PROVIDER:
        if request.get("reduction_percent") is not None:
            pct = Decimal(str(request["reduction_percent"]))
            payoff = money(billed * (Decimal("100") - pct) / Decimal("100"))
        elif request.get("use_ledger_current_due") and entry.get("reduction_entered"):
            payoff = money(entry["current_due"])
            pct = money((billed - payoff) / billed * Decimal("100"))
        else:
            raise NeedsInputError(f"No target reduction for {contact['name']}; ask the user for one")
        values.update(PAYOFF=fmt_money(payoff), REDUCTION_PERCENT=_pct(pct))
    else:
        if request.get("interest_reduction_percent") is None:
            raise NeedsInputError(f"No interest reduction percentage for {contact['name']}")
        pct = Decimal(str(request["interest_reduction_percent"]))
        payoff = money(entry["current_due"])
        good_through = to_date(entry.get("payoff_good_through"))
        values.update(PAYOFF=fmt_money(payoff), REDUCTION_PERCENT=_pct(pct),
                      PAYOFF_AS_OF=f", good through {good_through:%m/%d/%Y}" if good_through else "",
                      EXPIRED_NOTE="", PROPOSED_PAYOFF="")
        if good_through and good_through < as_of:
            flags.append(f"Payoff expired {good_through:%m/%d/%Y}; the letter asks for an updated payoff.")
            values["EXPIRED_NOTE"] = (" Because the payoff we have on file has expired, please "
                                      "include an updated payoff as of today.")
        if entry.get("principal") is not None and entry.get("interest") is not None:
            interest = money(entry["interest"])
            proposed = money(money(entry["principal"]) + interest * (Decimal("100") - pct) / Decimal("100"))
            values["PROPOSED_PAYOFF"] = f" Based on your figures, the revised payoff would be {fmt_money(proposed)}."
        else:
            flags.append("Principal and interest are not broken out; the letter requests the breakdown.")

    return Letter(entity=contact["name"], kind=kind, billed=billed, payoff=payoff,
                  reduction_percent=pct, filename=output_filename(contact["name"], client_name),
                  values=values, flags=flags)


def fill_template(template: Path, values: dict[str, str], out_path: Path) -> Path:
    """Replace placeholders run by run, so the template's formatting is kept."""
    doc = Document(template)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if _PLACEHOLDER.search(run.text):
                run.text = _PLACEHOLDER.sub(lambda m: values.get(m.group(0)[2:-2], m.group(0)), run.text)
    leftover = [p.text for p in doc.paragraphs if _PLACEHOLDER.search(p.text)]
    if leftover:
        raise ValueError(f"Unfilled placeholders in {out_path.name}: {leftover}")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def to_pdf(docx_path: Path, out_dir: Path) -> Path | None:
    """Convert with LibreOffice if it is installed. Returns None when it isn't."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
                   check=True, capture_output=True, timeout=120)
    pdf = out_dir / (docx_path.stem + ".pdf")
    return pdf if pdf.exists() else None


def check_pdf(pdf: Path, letter: Letter) -> list[str]:
    """QA on the final PDF: one page, and the key amounts actually appear in it."""
    try:
        from pypdf import PdfReader
    except ImportError:
        return []
    reader = PdfReader(pdf)
    text = " ".join(page.extract_text() or "" for page in reader.pages)
    problems = []
    if len(reader.pages) != 1:
        problems.append(f"{pdf.name} is {len(reader.pages)} pages; expected 1")
    for key in ("PAYOFF", "CLIENT_NAME", "ENTITY_NAME"):
        if letter.values[key] not in " ".join(text.split()):
            problems.append(f"{pdf.name} is missing {key.lower()} '{letter.values[key]}'")
    return problems


def draft_letters(claim: dict, contacts: list[dict], ledger: list[dict], request: dict,
                  out_dir: str | Path, keep_docx: bool = False) -> tuple[list[Letter], list[str]]:
    """Draft every requested letter. Returns (letters, notes) where notes cover holds and skips."""
    out = Path(out_dir)
    as_of = to_date(request.get("as_of")) or date.today()
    attorney = request.get("attorney_name", "[Attorney Name]")
    by_id = {c["contact_id"]: c for c in contacts}
    holds = set(request.get("hold") or [])
    notes = [f"Held, no letter drafted: {by_id[h]['name']}" for h in holds if h in by_id]

    work = out / "_work"
    templates = {k: build_template(k, work / f"{k}_template.docx") for k in (PROVIDER, LOAN)}
    letters: list[Letter] = []
    for item in request.get("requests") or []:
        contact_id = item["contact_id"]
        if contact_id in holds:
            notes.append(f"Skipped {by_id[contact_id]['name']}: it is on hold")
            continue
        try:
            letter = plan_letter(claim, by_id[contact_id], ledger, item, as_of, attorney)
        except NeedsInputError as exc:
            notes.append(f"Needs your input: {exc}")
            continue
        docx_path = fill_template(templates[letter.kind], letter.values,
                                  (out if keep_docx else work) / letter.filename.replace(".pdf", ".docx"))
        letter.docx_path = docx_path
        pdf = to_pdf(docx_path, out)
        if pdf:
            letter.pdf_path = pdf
            letter.flags += check_pdf(pdf, letter)
        else:
            notes.append(f"LibreOffice not found; kept {docx_path.name} instead of a PDF")
            if not keep_docx:
                letter.docx_path = Path(shutil.copy(docx_path, out / docx_path.name))
        letters.append(letter)

    if not keep_docx:
        shutil.rmtree(work, ignore_errors=True)
    return letters, notes


def summary(letters: list[Letter], notes: list[str]) -> str:
    lines = []
    for l in letters:
        what = (f"{l.reduction_percent.normalize():f}% reduction, payoff {fmt_money(l.payoff)}" if l.kind == PROVIDER
                else f"{l.reduction_percent.normalize():f}% interest reduction on payoff {fmt_money(l.payoff)}")
        lines.append(f"- {(l.pdf_path or l.docx_path).name}: {l.entity}, {what}")
        lines += [f"    - Flag: {f}" for f in l.flags]
    lines += [f"- {n}" for n in notes]
    return "\n".join(lines) + "\n"
