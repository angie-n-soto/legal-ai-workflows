"""Generic letter templates with {{PLACEHOLDERS}}.

The original skill edited the firm's approved Word templates. Those are not included here.
These stand-ins were written from scratch so the "fill the template, don't retype the
letter" approach can be shown. Every placeholder sits in its own run so filling it keeps
the run's font and formatting.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

FIRM_NAME = "Example Injury Law, PLLC"
FIRM_ADDRESS = "123 Placeholder Street, Suite 100  |  Dallas, TX 75201  |  (555) 010-0000"

PROVIDER_BODY = [
    "This firm represents {{CLIENT_NAME}} in a personal injury claim arising from a "
    "{{INCIDENT_TYPE}} on {{DOI}}. Our records show a balance of {{BILLED}} for treatment "
    "provided by your office.",
    "As we work to resolve this claim, we respectfully request that you reduce this balance "
    "by {{REDUCTION_PERCENT}}, to {{PAYOFF}}. A reduction helps our client receive a fair net "
    "recovery, and payment will be issued promptly once settlement funds are disbursed.",
    "If you agree, please sign the acceptance below and return it to our office.",
]
PROVIDER_ACCEPTANCE = (
    "{{ENTITY_NAME}} agrees to accept {{PAYOFF}} as payment in full for services provided "
    "to {{CLIENT_NAME}}."
)

LOAN_BODY = [
    "This firm represents {{CLIENT_NAME}} in a personal injury claim arising from a "
    "{{INCIDENT_TYPE}} on {{DOI}}. Our records reflect a current payoff of {{PAYOFF}} for the "
    "funding your company provided to our client{{PAYOFF_AS_OF}}.",
    "We respectfully request a {{REDUCTION_PERCENT}} reduction of the interest charges "
    "associated with this funding. This request applies to interest only, not principal. "
    "The reduction would help maximize our client's net recovery and allow this matter to "
    "be resolved promptly.{{PROPOSED_PAYOFF}}",
    "Please provide a breakdown of principal and accrued interest, together with a revised "
    "payoff that reflects the requested interest reduction.{{EXPIRED_NOTE}}",
    "Please confirm any approved reduction and the updated payoff in writing.",
]
LOAN_ACCEPTANCE = (
    "{{ENTITY_NAME}} agrees to reduce the interest charges on {{CLIENT_NAME}}'s account by "
    "{{REDUCTION_PERCENT}}. Revised payoff: $______________"
)


def _add(paragraph, text: str, bold: bool = False, size: int = 11):
    """Add text to a paragraph, putting each {{PLACEHOLDER}} in its own run."""
    import re

    for piece in re.split(r"(\{\{[A-Z_]+\}\})", text):
        if piece:
            run = paragraph.add_run(piece)
            run.bold = bold
            run.font.size = Pt(size)
    return paragraph


def build_template(kind: str, path: str | Path) -> Path:
    body, acceptance = (PROVIDER_BODY, PROVIDER_ACCEPTANCE) if kind == "provider" else (LOAN_BODY, LOAN_ACCEPTANCE)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Inches(1))
    section.top_margin = section.bottom_margin = Inches(0.8)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)

    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = head.add_run(FIRM_NAME)
    run.bold, run.font.size = True, Pt(16)
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(FIRM_ADDRESS).font.size = Pt(9)

    _add(doc.add_paragraph(), "{{LETTER_DATE}}")
    _add(doc.add_paragraph(), "{{ENTITY_NAME}}", bold=True)
    _add(doc.add_paragraph(), "{{ENTITY_ADDRESS}}")
    _add(doc.add_paragraph(), "RE:\tOur Client: {{CLIENT_NAME}}\n\tDate of Birth: {{DOB}}\n\tDate of Incident: {{DOI}}", bold=True)
    _add(doc.add_paragraph(), "Dear Billing Department:" if kind == "provider" else "To Whom It May Concern:")
    for text in body:
        _add(doc.add_paragraph(), text)
    _add(doc.add_paragraph(), "Sincerely,")
    _add(doc.add_paragraph(), "[Signature]\n{{ATTORNEY_NAME}}\n" + FIRM_NAME)

    doc.add_paragraph("_" * 80)
    _add(doc.add_paragraph(), "ACCEPTANCE", bold=True)
    _add(doc.add_paragraph(), acceptance)
    _add(doc.add_paragraph(), "Print Name: ______________________________")
    _add(doc.add_paragraph(), "Signature: _______________________________     Date: ______________")

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
